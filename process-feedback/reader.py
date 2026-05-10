#!/usr/bin/env python

import os
import pickle
import openpyxl
import anytree


class Category(anytree.NodeMixin):

    def __init__(self, name, id, parent_id, path, export_path='.', parent=None, children=None):
        super().__init__()
        self.name = name
        self.id = id
        self.parent_id = parent_id
        self.category_path = path
        self.parent = parent
        if children:
            self.children = children
        self.result = {}
        self.export_path = export_path

    def set_parent(self, parent):
        self.parent = parent

    def has_parent(self):
        return self.parent

    def set_export_path(self, export_path):
        self.export_path = export_path

    def set_result(self, result):
        if result['Licenta']['courses']:
            self.result['Licenta'] = result['Licenta']
        else:
            self.result['Licenta'] = None
        if result['Masterat']['courses']:
            self.result['Masterat'] = result['Masterat']
        else:
            self.result['Masterat'] = None
        if self.result['Licenta'] and self.result['Masterat']:
            self.result['Cumulat'] = result['Cumulat']
        else:
            self.result['Cumulat'] = None

    def print_result_line(data):
        value = ""
        for k, v in data.items():
            if not v['value']:
                value += ',  None'
                continue
            if k == 'name':
                value += '{:40s}'.format(v['value'])
            elif k == 'num':
                value += '{:6d}'.format(v['value'])
            elif k == 'percentage':
                value += ',{:5.2f}%'.format(v['value'])
            else:
                value += ',{:6.2f}'.format(v['value'])
        print(value)

    def print_result_part(result):
        print("\n= overall\n")
        Category.print_result_line(result['overall'])
        print("\n= courses\n")
        for k, i in result['courses'].items():
            Category.print_result_line(i)
        print("\n= profs\n")
        for k, i in result['profs'].items():
            Category.print_result_line(i)
        print("\n= assists\n")
        for k, i in result['assists'].items():
            Category.print_result_line(i)
        print("rejected: {}".format(len(result['rejected'])))
        print("blacklisted: {}".format(len(result['blacklisted'])))
        print("courses with feedback: {}".format(len(result['courses'])))
        print("courses with no feedback: {}".format(len(result['courses_list']) - len(result['courses'])))

    def print_result(self):
        if self.result['Cumulat']:
            print("\n==== CUMULAT ====\n")
            Category.print_result_part(self.result['Cumulat'])
        if self.result['Licenta']:
            print("\n==== LICENȚĂ ====\n")
            Category.print_result_part(self.result['Licenta'])
        if self.result['Masterat']:
            print("\n==== MASTERAT ====\n")
            Category.print_result_part(self.result['Masterat'])

    def write_line_in_overall(ws, key, idx, result):
        ws['A{:d}'.format(idx)] = key
        ws['B{:d}'.format(idx)] = len(result['courses'])
        ws['C{:d}'.format(idx)] = len(result['courses_list']) - len(result['courses'])
        ws['D{:d}'.format(idx)] = result['overall']['num']['value']
        ws['E{:d}'.format(idx)] = result['overall']['percentage']['value'] / 100
        ws['E{:d}'.format(idx)].number_format = openpyxl.styles.numbers.FORMAT_PERCENTAGE_00
        ws['F{:d}'.format(idx)] = result['overall']['eval_overall']['value']
        ws['F{:d}'.format(idx)].number_format = openpyxl.styles.numbers.FORMAT_NUMBER_00
        ws['G{:d}'.format(idx)] = result['overall']['eval_overall']['stdev']
        ws['G{:d}'.format(idx)].number_format = openpyxl.styles.numbers.FORMAT_NUMBER_00

    def export_overall(self, ws, part, maxlevel):
        for c_id in range(ord('A'), ord('H')):
            ws.column_dimensions[""+chr(c_id)].width = 20
        ws.row_dimensions[1].height = 36
        for c_id in range(1, ord('I') - ord('A')):
            ws.cell(row=1, column=c_id).alignment = openpyxl.styles.Alignment(wrap_text=True, vertical='top')
            ws.cell(row=1, column=c_id).font = openpyxl.styles.Font(bold=True)
        ws['A1'] = "Grup"
        ws['B1'] = "Cursuri cu feedback"
        ws['C1'] = "Cursuri fără feedback"
        ws['D1'] = "Număr de feedbackuri"
        ws['E1'] = "Procentaj de completare feedback"
        ws['F1'] = "Evaluare generală discipline (medie)"
        ws['G1'] = "Evaluare generală discipline (abatere standard)"

        idx = 2
        for node in anytree.PreOrderIter(self, maxlevel=maxlevel):
            if part == 'Cumulat':
                if node.result['Cumulat']:
                    Category.write_line_in_overall(ws, '{} (Cumulat)'.format(node.name), idx, node.result['Cumulat'])
                    idx += 1
                if node.result['Licenta']:
                    Category.write_line_in_overall(ws, '{} (Licenta)'.format(node.name), idx, node.result['Licenta'])
                    idx += 1
                if node.result['Masterat']:
                    Category.write_line_in_overall(ws, '{} (Masterat)'.format(node.name), idx, node.result['Masterat'])
                    idx += 1
            else:
                if node.result[part]:
                    Category.write_line_in_overall(ws, '{} ({})'.format(node.name, part), idx, node.result[part])
                    idx += 1

    def export_full_for_component(self, ws, part, component):
        for c_id in range(ord('A'), ord('Z')):
            ws.column_dimensions[""+chr(c_id)].width = 20
        ws.row_dimensions[1].height = 60
        for column in range(1, ord('Z') - ord('A')):
            ws.cell(row=1, column=column).alignment = openpyxl.styles.Alignment(wrap_text=True, vertical='top')
            ws.cell(row=1, column=column).font = openpyxl.styles.Font(bold=True)

        if not self.result[part][component]:
            return

        k = next(iter(self.result[part][component]))
        v = self.result[part][component][k]
        i = 0
        for k, item in v.items():
            ws['{:c}1'.format((ord('A')+i))] = item['title']
            i += 1

        row = 2
        for k, v in self.result[part][component].items():
            column = 1
            for k2, item in v.items():
                cell = ws.cell(row=row, column=column)
                if item['value'] == None:
                    cell.value = 'N/A'
                else:
                    cell.value = item['value']
                    if column == 4:
                        cell.number_format = openpyxl.styles.numbers.FORMAT_PERCENTAGE_00
                        cell.value = item['value'] / 100
                    if column > 4:
                        cell.number_format = openpyxl.styles.numbers.FORMAT_NUMBER_00
                column += 1
            row += 1

    def export_courses_summary(self, ws, part):
        for c_id in range(ord('A'), ord('F')):
            ws.column_dimensions[""+chr(c_id)].width = 20
        ws.row_dimensions[1].height = 40
        for column in range(1, ord('G') - ord('A')):
            ws.cell(row=1, column=column).alignment = openpyxl.styles.Alignment(wrap_text=True, vertical='top')
            ws.cell(row=1, column=column).font = openpyxl.styles.Font(bold=True)

        comp = self.result[part]['courses']
        if not comp:
            return

        k = next(iter(comp))
        v = comp[k]
        ws['A1'] = v['name']['title']
        ws['B1'] = v['num']['title']
        ws['C1'] = v['percentage']['title']
        ws['D1'] = "Evaluare generală discipline (medie)"
        ws['E1'] = "Evaluare generală discipline (abatere standard)"

        row = 2
        for k, v in comp.items():
            ws['A{:d}'.format(row)] = v['name']['value']
            ws['B{:d}'.format(row)] = v['num']['value']
            ws['C{:d}'.format(row)] = v['percentage']['value'] / 100
            ws['C{:d}'.format(row)].number_format = openpyxl.styles.numbers.FORMAT_PERCENTAGE_00
            ws['D{:d}'.format(row)] = v['eval_overall']['value']
            ws['D{:d}'.format(row)].number_format = openpyxl.styles.numbers.FORMAT_NUMBER_00
            ws['E{:d}'.format(row)] = v['eval_overall']['stdev']
            ws['E{:d}'.format(row)].number_format = openpyxl.styles.numbers.FORMAT_NUMBER_00
            row += 1

    def export_profs_summary(self, ws, part):
        for c_id in range(ord('A'), ord('F')):
            ws.column_dimensions[""+chr(c_id)].width = 20
        ws.row_dimensions[1].height = 40
        for column in range(1, ord('G') - ord('A')):
            ws.cell(row=1, column=column).alignment = openpyxl.styles.Alignment(wrap_text=True, vertical='top')
            ws.cell(row=1, column=column).font = openpyxl.styles.Font(bold=True)

        comp = self.result[part]['profs']
        if not comp:
            return

        k = next(iter(comp))
        v = comp[k]
        ws['A1'] = v['name']['title']
        ws['B1'] = v['num']['title']
        ws['C1'] = v['percentage']['title']
        ws['D1'] = "Evaluare agregată titular curs (medie)"
        ws['E1'] = "Evaluare agregată titular curs (abatere standard)"

        row = 2
        for k, v in comp.items():
            ws['A{:d}'.format(row)] = v['name']['value']
            ws['B{:d}'.format(row)] = v['num']['value']
            ws['C{:d}'.format(row)] = v['percentage']['value'] / 100
            ws['C{:d}'.format(row)].number_format = openpyxl.styles.numbers.FORMAT_PERCENTAGE_00
            ws['D{:d}'.format(row)] = v['overall_prof']['value']
            ws['D{:d}'.format(row)].number_format = openpyxl.styles.numbers.FORMAT_NUMBER_00
            ws['E{:d}'.format(row)] = v['overall_prof']['stdev']
            ws['E{:d}'.format(row)].number_format = openpyxl.styles.numbers.FORMAT_NUMBER_00
            row += 1

    def export_assists_summary(self, ws, part):
        for c_id in range(ord('A'), ord('F')):
            ws.column_dimensions[""+chr(c_id)].width = 20
        ws.row_dimensions[1].height = 40
        for column in range(1, ord('G') - ord('A')):
            ws.cell(row=1, column=column).alignment = openpyxl.styles.Alignment(wrap_text=True, vertical='top')
            ws.cell(row=1, column=column).font = openpyxl.styles.Font(bold=True)

        comp = self.result[part]['assists']
        if not comp:
            return

        k = next(iter(comp))
        v = comp[k]
        ws['A1'] = v['name']['title']
        ws['B1'] = v['num']['title']
        ws['C1'] = "Evaluare agregată titular laborator (medie)"
        ws['D1'] = "Evaluare agregată titular laborator (abatere standard)"

        row = 2
        for k, v in comp.items():
            ws['A{:d}'.format(row)] = v['name']['value']
            ws['B{:d}'.format(row)] = v['num']['value']
            ws['C{:d}'.format(row)] = v['overall_prof']['value']
            ws['C{:d}'.format(row)].number_format = openpyxl.styles.numbers.FORMAT_NUMBER_00
            ws['D{:d}'.format(row)] = v['overall_prof']['stdev']
            ws['D{:d}'.format(row)].number_format = openpyxl.styles.numbers.FORMAT_NUMBER_00
            row += 1

    def export_zone_courses(self, ws, part):
        for c_id in range(ord('A'), ord('E')):
            ws.column_dimensions[""+chr(c_id)].width = 20
        ws.row_dimensions[1].height = 25
        for column in range(1, ord('F') - ord('A')):
            ws.cell(row=1, column=column).alignment = openpyxl.styles.Alignment(wrap_text=True, vertical='top')
            ws.cell(row=1, column=column).font = openpyxl.styles.Font(bold=True)

        comp = self.result[part]['courses']
        if not comp:
            return

        ws['A1'] = 'Zona / Cvartilă'
        ws['B1'] = 'Număr de cursuri'
        ws['C1'] = 'Procent de cursuri'

        for i in range(1, 5):
            num = len([v for k, v in comp.items() if v['eval_overall']['value'] >= i and v['eval_overall']['value'] < i+1])
            if i == 4:
                num += len([v for k, v in comp.items() if v['eval_overall']['value'] == 5])
            all_num = len(comp)
            ws['A{}'.format(i+1)] = '{}-{}'.format(i, i+1)
            ws['B{}'.format(i+1)] = num
            if all_num == 0:
                perc = 0.0
            else:
                perc = float(num) / all_num
            ws['C{}'.format(i+1)] = perc
            ws['C{}'.format(i+1)].number_format = openpyxl.styles.numbers.FORMAT_PERCENTAGE_00

    def export_zone_profs(self, ws, part):
        for c_id in range(ord('A'), ord('E')):
            ws.column_dimensions[""+chr(c_id)].width = 20
        ws.row_dimensions[1].height = 25
        for column in range(1, ord('F') - ord('A')):
            ws.cell(row=1, column=column).alignment = openpyxl.styles.Alignment(wrap_text=True, vertical='top')
            ws.cell(row=1, column=column).font = openpyxl.styles.Font(bold=True)

        comp = self.result[part]['profs']
        if not comp:
            return

        ws['A1'] = 'Zona / Cvartilă'
        ws['B1'] = 'Număr de titulari de curs'
        ws['C1'] = 'Procent de titulari de curs'

        for i in range(1, 5):
            num = len([v for k, v in comp.items() if v['overall_prof']['value'] >= i and v['overall_prof']['value'] < i+1])
            if i == 4:
                num += len([v for k, v in comp.items() if v['overall_prof']['value'] == 5])
            all_num = len(comp)
            ws['A{}'.format(i+1)] = '{}-{}'.format(i, i+1)
            ws['B{}'.format(i+1)] = num
            if all_num == 0:
                perc = 0.0
            else:
                perc = float(num) / all_num
            ws['C{}'.format(i+1)] = perc
            ws['C{}'.format(i+1)].number_format = openpyxl.styles.numbers.FORMAT_PERCENTAGE_00

    def export_zone_assists(self, ws, part):
        for c_id in range(ord('A'), ord('E')):
            ws.column_dimensions[""+chr(c_id)].width = 20
        ws.row_dimensions[1].height = 25
        for column in range(1, ord('F') - ord('A')):
            ws.cell(row=1, column=column).alignment = openpyxl.styles.Alignment(wrap_text=True, vertical='top')
            ws.cell(row=1, column=column).font = openpyxl.styles.Font(bold=True)

        comp = self.result[part]['assists']
        if not comp:
            return

        ws['A1'] = 'Zona / Cvartilă'
        ws['B1'] = 'Număr de titulari de laborator'
        ws['C1'] = 'Procent de titulari de laborator'

        for i in range(1, 5):
            num = len([v for k, v in comp.items() if v['overall_assist']['value'] >= i and v['overall_assist']['value'] < i+1])
            if i == 4:
                num += len([v for k, v in comp.items() if v['overall_assist']['value'] == 5])
            all_num = len(comp)
            ws['A{}'.format(i+1)] = '{}-{}'.format(i, i+1)
            ws['B{}'.format(i+1)] = num
            if all_num == 0:
                perc = 0.0
            else:
                perc = float(num) / all_num
            ws['C{}'.format(i+1)] = perc
            ws['C{}'.format(i+1)].number_format = openpyxl.styles.numbers.FORMAT_PERCENTAGE_00

    def export_spreadsheet_part(self, part):
        wb = openpyxl.Workbook()

        ws = wb.active
        ws.title = "Ansamblu"
        self.export_overall(ws, part, 2)

        wb.create_sheet("Ansamblu (nivel 2)")
        ws = wb["Ansamblu (nivel 2)"]
        self.export_overall(ws, part, 3)

        wb.create_sheet("Cursuri (sumar)")
        ws = wb["Cursuri (sumar)"]
        self.export_courses_summary(ws, part)

        wb.create_sheet("Titulari curs (sumar)")
        ws = wb["Titulari curs (sumar)"]
        self.export_profs_summary(ws, part)

        wb.create_sheet("Titulari laborator (sumar)")
        ws = wb["Titulari laborator (sumar)"]
        self.export_assists_summary(ws, part)

        wb.create_sheet("Cursuri (zone)")
        ws = wb["Cursuri (zone)"]
        self.export_zone_courses(ws, part)

        wb.create_sheet("Titulari (zone)")
        ws = wb["Titulari (zone)"]
        self.export_zone_profs(ws, part)

        wb.create_sheet("Titulari laborator (zone)")
        ws = wb["Titulari laborator (zone)"]
        self.export_zone_assists(ws, part)

        wb.create_sheet("Cursuri (complet)")
        ws = wb["Cursuri (complet)"]
        self.export_full_for_component(ws, part, 'courses')

        wb.create_sheet("Titulari curs (complet)")
        ws = wb["Titulari curs (complet)"]
        self.export_full_for_component(ws, part, 'profs')

        wb.create_sheet("Titulari laborator (complet)")
        ws = wb["Titulari laborator (complet)"]
        self.export_full_for_component(ws, part, 'assists')

        wb.save(os.path.join(self.export_path, '{}.xlsx'.format(part)))

    def export_spreadsheet(self):
        print("Exporting category {} to {}".format(self.name, self.export_path))
        for k in self.result:
            if self.result[k]:
                if not os.path.exists(self.export_path):
                    os.mkdir(self.export_path)
                self.export_spreadsheet_part(k)


class Reader():

    def __init__(self, courses_file, categories_file, processed_courses_dir, processed_categories_dir):
        self.courses = pickle.load(open(courses_file, 'rb'))
        self.categories = pickle.load(open(categories_file, 'rb'))
        self.processed_courses_dir = processed_courses_dir
        self.processed_categories_dir = processed_categories_dir
        self.root = None
        self.category_list = []

    def build_category(self, category_id):
        c = next((c for c in self.categories if c['id'] == category_id))
        c = Category(c['name'], c['id'], c['parent'], c['path'])

        r = pickle.load(open(os.path.join(self.processed_categories_dir, '{}.p'.format(c.id)), "rb"))
        c.set_result({
            'Cumulat': pickle.load(open(os.path.join(self.processed_categories_dir, '{}.p'.format(c.id)), "rb")),
            'Licenta': pickle.load(open(os.path.join(self.processed_categories_dir, '{}_bachelor.p'.format(c.id)), "rb")),
            'Masterat': pickle.load(open(os.path.join(self.processed_categories_dir, '{}_master.p'.format(c.id)), "rb"))
            })

        return c

    def build_all_categories(self):
        for c in self.categories:
            c = Category(c['name'], c['id'], c['parent'], c['path'])

            r = pickle.load(open(os.path.join(self.processed_categories_dir, '{}.p'.format(c.id)), "rb"))
            c.set_result({
                'Cumulat': pickle.load(open(os.path.join(self.processed_categories_dir, '{}.p'.format(c.id)), "rb")),
                'Licenta': pickle.load(open(os.path.join(self.processed_categories_dir, '{}_bachelor.p'.format(c.id)), "rb")),
                'Masterat': pickle.load(open(os.path.join(self.processed_categories_dir, '{}_master.p'.format(c.id)), "rb"))
                })
            self.category_list.append(c)
        self.root = next((c for c in self.category_list if c.id == 1))
        self.root.name = 'UPB'
        self.root.parent_id = -1

        for c in self.category_list:
            if c.parent_id == 0:
                c.set_parent(self.root)
                continue
            for c2 in self.category_list:
                if c2.id == c.parent_id:
                    c.set_parent(c2)

    def export_category(self, category_id):
        c = next((c for c in self.category_list if c.id == category_id))
        c.export_spreadsheet()

    def export_all_categories(self, export_dir):
        for c in anytree.PreOrderIter(self.root):
            if c.has_parent():
                c.set_export_path(os.path.join(c.parent.export_path, c.name))
            else:
                c.set_export_path(os.path.join(export_dir, c.name))
            c.export_spreadsheet()

    def print_all_categories(self):
        for pre, fill, node in anytree.RenderTree(self.root):
            treestr = u"{}{} ({})".format(pre, node.name, node.id)
            print(treestr.ljust(8))
